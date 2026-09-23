import streamlit as st
from datetime import date
from sqlalchemy import create_engine, Column, Integer, String, Text, Date, ForeignKey
from sqlalchemy.orm import declarative_base, sessionmaker

st.set_page_config(page_title="Facilita Manager", page_icon="🏛️", layout="wide")
# =========================================================
# IDENTIDADE VISUAL DA FACILITA
# =========================================================

st.markdown(
    """
    <style>
        /* Fundo geral da aplicação */
        .stApp {
            background-color: #f4f7fc;
        }

        /* Menu lateral */
        section[data-testid="stSidebar"] {
            background: linear-gradient(
                180deg,
                #071d52 0%,
                #0b2d78 100%
            );
        }

        section[data-testid="stSidebar"] * {
            color: white;
        }

        /* Área principal */
        .main .block-container {
            padding-top: 2rem;
            padding-bottom: 2rem;
            max-width: 1500px;
        }

        /* Títulos */
        h1, h2, h3 {
            color: #092b72 !important;
        }

        /* Cartões de métricas */
        div[data-testid="stMetric"] {
            background-color: white;
            border: 1px solid #e1e8f2;
            border-radius: 14px;
            padding: 18px;
            box-shadow: 0 3px 12px rgba(9, 43, 114, 0.06);
        }

        div[data-testid="stMetricLabel"] {
            color: #52627a !important;
        }

        div[data-testid="stMetricValue"] {
            color: #092b72 !important;
            font-weight: 700;
        }

        /* Botões */
        .stButton > button {
            background-color: #092b72;
            color: white;
            border: none;
            border-radius: 8px;
            min-height: 42px;
            font-weight: 600;
        }

        .stButton > button:hover {
            background-color: #f4b400;
            color: #092b72;
            border: none;
        }

        /* Campos de formulário */
        div[data-baseweb="input"] > div,
        div[data-baseweb="select"] > div,
        textarea {
            border-radius: 8px !important;
        }

        /* Abas */
        button[data-baseweb="tab"] {
            font-weight: 600;
        }

        /* Expanders */
        div[data-testid="stExpander"] {
            background-color: white;
            border: 1px solid #e1e8f2;
            border-radius: 10px;
        }

        /* Linha divisória */
        hr {
            border-color: #dce5f2;
        }

        /* Texto auxiliar */
        .texto-institucional {
            color: #64748b;
            font-size: 0.95rem;
        }

        /* Cartões personalizados */
        .cartao-facilita {
            background-color: white;
            border: 1px solid #e1e8f2;
            border-radius: 14px;
            padding: 20px;
            box-shadow: 0 3px 12px rgba(9, 43, 114, 0.06);
        }
     /* =====================================================
   CORREÇÃO DE CONTRASTE
   ===================================================== */

/* Textos dentro dos cartões */
div[data-testid="stVerticalBlockBorderWrapper"] p {
    color: #334155 !important;
}

/* Textos gerais da área principal */
.main p {
    color: #334155;
}

/* Textos dos expanders */
div[data-testid="stExpander"] p {
    color: #334155 !important;
}

/* Títulos dentro dos cartões */
div[data-testid="stVerticalBlockBorderWrapper"] h1,
div[data-testid="stVerticalBlockBorderWrapper"] h2,
div[data-testid="stVerticalBlockBorderWrapper"] h3 {
    color: #092b72 !important;
}

/* Métricas do Dashboard */
div[data-testid="stMetric"] label {
    color: #52627a !important;
}

div[data-testid="stMetric"] div {
    color: #092b72 !important;
}

/* Textos dos campos */
label {
    color: #334155 !important;
}   
    </style>
    """,
    unsafe_allow_html=True,
)
engine = create_engine("sqlite:///facilita.db", connect_args={"check_same_thread": False})
Base = declarative_base()
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)


class Cliente(Base):
	__tablename__ = "clientes"
	id = Column(Integer, primary_key=True, index=True)
	nome = Column(String(150), nullable=False)
	tipo = Column(String(50), nullable=False)
	cidade = Column(String(100), nullable=False)
	estado = Column(String(2), nullable=False)
	responsavel = Column(String(100), nullable=False)
	email = Column(String(150))
	telefone = Column(String(30))
	observacoes = Column(Text)


class Demanda(Base):
	__tablename__ = "demandas"
	id = Column(Integer, primary_key=True, index=True)
	cliente_id = Column(Integer, ForeignKey("clientes.id"), nullable=False)
	titulo = Column(String(200), nullable=False)
	descricao = Column(Text)
	responsavel = Column(String(100), nullable=False)
	status = Column(String(50), nullable=False)
	prioridade = Column(String(30), nullable=False)
	prazo = Column(Date)
	data_criacao = Column(Date, nullable=False)
	observacoes = Column(Text)


Base.metadata.create_all(bind=engine)
RESPONSAVEIS = ["Vitor", "Ítalo", "Jadsson", "Álvaro Vinícius", "Não definido"]
STATUS = ["A Fazer", "Em Andamento", "Aguardando Cliente", "Concluído"]


def clientes():
	with SessionLocal() as db:
		return db.query(Cliente).order_by(Cliente.nome).all()


def demandas():
	with SessionLocal() as db:
		rows = db.query(Demanda).order_by(Demanda.id.desc()).all()
		return [(d, db.query(Cliente).filter(Cliente.id == d.cliente_id).first()) for d in rows]


def excluir(model, ident):
	with SessionLocal() as db:
		item = db.query(model).filter(model.id == ident).first()
		if item:
			db.delete(item)
			db.commit()
def importar_carteira_excel(caminho_arquivo):
    from openpyxl import load_workbook

    wb = load_workbook(caminho_arquivo, data_only=True)
    ws = wb["CARTEIRA"]

    importados = 0
    ignorados = 0

    colunas = {
        "Vitor": {
            "camara": 2,
            "prefeitura": 3,
        },
        "Ítalo": {
            "camara": 5,
            "prefeitura": 6,
        },
        "Jadsson": {
            "camara": 8,
            "prefeitura": 9,
        },
    }

    with SessionLocal() as db:
        for responsavel, tipos in colunas.items():

            for tipo_nome, coluna in [
                ("Câmara Municipal", tipos["camara"]),
                ("Prefeitura", tipos["prefeitura"]),
            ]:

                for linha in range(4, ws.max_row + 1):
                    valor = ws.cell(linha, coluna).value

                    if not valor:
                        continue

                    nome = str(valor).strip()

                    existente = (
                        db.query(Cliente)
                        .filter(Cliente.nome == nome)
                        .first()
                    )

                    if existente:
                        # Se o cliente já existe, apenas corrige
                        # o responsável e o tipo.
                        existente.responsavel = responsavel
                        existente.tipo = tipo_nome
                        ignorados += 1
                        continue

                    novo_cliente = Cliente(
                        nome=nome,
                        tipo=tipo_nome,
                        cidade="Não informada",
                        estado="SE",
                        responsavel=responsavel,
                        email="",
                        telefone="",
                        observacoes="Importado da carteira 2026."
                    )

                    db.add(novo_cliente)
                    importados += 1

        db.commit()

    return importados, ignorados


st.sidebar.image(
    "assets/logo-facilita.png",
    use_container_width=True,
)

pagina = st.sidebar.radio(
    "Navegação",
    [
        "Dashboard",
        "Clientes",
        "Demandas",
        "Kanban",
        "Carteiras",
        "Relatórios",
        "Configurações",
    ],
)
st.sidebar.markdown(
    """
    <div style="
        text-align: center;
        color: #f4b400;
        font-size: 12px;
        margin-top: -8px;
        margin-bottom: 18px;
    ">
        Sistema de Gestão da Facilita
    </div>
    """,
    unsafe_allow_html=True,
)

st.sidebar.markdown("---")
if pagina == "Dashboard":
	cs, ds = clientes(), demandas()
	abertos = [x for x in ds if x[0].status != "Concluído"]
	st.title("📊 Dashboard")
	a, b, c, d = st.columns(4)
	a.metric("Clientes cadastrados", len(cs))
	b.metric("Demandas abertas", len(abertos))
	c.metric("Demandas concluídas", len(ds) - len(abertos))
	d.metric("Responsáveis", len(set(x.responsavel for x in cs)))
	st.subheader("Demandas em aberto")
	if not abertos:
		st.info("Não existem demandas abertas.")
	for demanda, cliente in abertos[:5]:
		st.write(f"**{demanda.titulo}** — {cliente.nome if cliente else 'Cliente não encontrado'} — {demanda.status}")

elif pagina == "Clientes":
	st.title("👥 Clientes")
	cadastro, lista = st.tabs(["➕ Cadastrar cliente", "📋 Clientes cadastrados"])
	with cadastro:
		with st.form("cliente"):
			nome = st.text_input("Nome do órgão público*")
			tipo = st.selectbox("Tipo de órgão*", ["Prefeitura", "Câmara Municipal", "Outro"])
			cidade, estado = st.text_input("Cidade*"), st.selectbox("Estado*", ["SE", "BA", "AL", "PE", "PB", "CE", "PI", "MA", "RN", "Outro"])
			resp = st.selectbox("Responsável pela carteira*", RESPONSAVEIS)
			email, telefone = st.text_input("E-mail"), st.text_input("Telefone ou WhatsApp")
			obs = st.text_area("Observações")
			if st.form_submit_button("💾 Salvar cliente", use_container_width=True):
				if not nome.strip() or not cidade.strip():
					st.error("Informe o nome e a cidade.")
				else:
					with SessionLocal() as db:
						db.add(Cliente(nome=nome.strip(), tipo=tipo, cidade=cidade.strip(), estado=estado, responsavel=resp, email=email.strip(), telefone=telefone.strip(), observacoes=obs.strip()))
						db.commit()
					st.success("Cliente cadastrado com sucesso!")
					st.rerun()
	with lista:
		for cliente in clientes():
			with st.expander(f"{cliente.nome} — {cliente.cidade}/{cliente.estado}"):
				st.write(f"**Tipo:** {cliente.tipo}  |  **Responsável:** {cliente.responsavel}")
				st.write(f"**E-mail:** {cliente.email or 'Não informado'}  |  **Telefone:** {cliente.telefone or 'Não informado'}")
				if cliente.observacoes: st.write(f"**Observações:** {cliente.observacoes}")
				if st.button("🗑️ Excluir cliente", key=f"ec{cliente.id}"):
					excluir(Cliente, cliente.id); st.rerun()

elif pagina == "Demandas":
    st.title("📋 Demandas")

    cs = clientes()

    nova, lista = st.tabs([
        "➕ Nova demanda",
        "📋 Demandas cadastradas"
    ])

    with nova:
        if not cs:
            st.warning(
                "Cadastre pelo menos um cliente antes de criar uma demanda."
            )
        else:
            with st.form("demanda"):
                mapa = {
                    f"{c.nome} — {c.cidade}/{c.estado}": c.id
                    for c in cs
                }

                cliente = st.selectbox(
                    "Cliente*",
                    list(mapa.keys())
                )

                titulo = st.text_input(
                    "Título da demanda*"
                )

                descricao = st.text_area(
                    "Descrição da solicitação"
                )

                responsavel = st.selectbox(
                    "Responsável*",
                    RESPONSAVEIS[:-1]
                )

                status = st.selectbox(
                    "Status*",
                    STATUS
                )

                prioridade = st.selectbox(
                    "Prioridade*",
                    ["Alta", "Média", "Baixa"]
                )

                prazo = st.date_input(
                    "Prazo",
                    value=date.today()
                )

                obs = st.text_area(
                    "Observações internas"
                )

                salvar = st.form_submit_button(
                    "💾 Salvar demanda",
                    use_container_width=True
                )

                if salvar:
                    if not titulo.strip():
                        st.error(
                            "Informe o título da demanda."
                        )
                    else:
                        with SessionLocal() as db:
                            db.add(
                                Demanda(
                                    cliente_id=mapa[cliente],
                                    titulo=titulo.strip(),
                                    descricao=descricao.strip(),
                                    responsavel=responsavel,
                                    status=status,
                                    prioridade=prioridade,
                                    prazo=prazo,
                                    data_criacao=date.today(),
                                    observacoes=obs.strip()
                                )
                            )

                            db.commit()

                        st.success(
                            "Demanda cadastrada com sucesso!"
                        )

                        st.rerun()

    with lista:
        st.subheader("🔎 Filtros")

        col1, col2, col3, col4 = st.columns(4)

        with col1:
            filtro_status = st.selectbox(
                "Status",
                ["Todos"] + STATUS
            )

        with col2:
            filtro_responsavel = st.selectbox(
                "Responsável",
                ["Todos"] + RESPONSAVEIS[:-1]
            )

        with col3:
            filtro_prioridade = st.selectbox(
                "Prioridade",
                ["Todas", "Alta", "Média", "Baixa"]
            )

        with col4:
            opcoes_clientes = ["Todos"] + [
                f"{c.nome} — {c.cidade}/{c.estado}"
                for c in cs
            ]

            filtro_cliente = st.selectbox(
                "Cliente",
                opcoes_clientes
            )

        st.divider()

        resultados = []

        for demanda, cliente in demandas():

            if (
                filtro_status != "Todos"
                and demanda.status != filtro_status
            ):
                continue

            if (
                filtro_responsavel != "Todos"
                and demanda.responsavel != filtro_responsavel
            ):
                continue

            if (
                filtro_prioridade != "Todas"
                and demanda.prioridade != filtro_prioridade
            ):
                continue

            if filtro_cliente != "Todos":
                nome_cliente = (
                    f"{cliente.nome} — "
                    f"{cliente.cidade}/{cliente.estado}"
                    if cliente
                    else "Cliente não encontrado"
                )

                if nome_cliente != filtro_cliente:
                    continue

            resultados.append(
                (demanda, cliente)
            )

        st.write(
            f"**{len(resultados)} demanda(s) encontrada(s).**"
        )

        if not resultados:
            st.info(
                "Nenhuma demanda encontrada com esses filtros."
            )

        for demanda, cliente in resultados:

            nome_cliente = (
                cliente.nome
                if cliente
                else "Cliente não encontrado"
            )

            with st.expander(
                f"{demanda.titulo} — {nome_cliente}"
            ):

                st.write(
                    f"**Responsável:** "
                    f"{demanda.responsavel}"
                )

                st.write(
                    f"**Status:** "
                    f"{demanda.status}"
                )

                st.write(
                    f"**Prioridade:** "
                    f"{demanda.prioridade}"
                )

                if demanda.prazo:
                    st.write(
                        f"**Prazo:** "
                        f"{demanda.prazo.strftime('%d/%m/%Y')}"
                    )
                else:
                    st.write(
                        "**Prazo:** Não informado"
                    )

                if demanda.descricao:
                    st.write(
                        f"**Descrição:** "
                        f"{demanda.descricao}"
                    )

                if demanda.observacoes:
                    st.write(
                        f"**Observações:** "
                        f"{demanda.observacoes}"
                    )

                novo = st.selectbox(
                    "Alterar status",
                    STATUS,
                    index=STATUS.index(demanda.status),
                    key=f"s{demanda.id}"
                )

                if st.button(
                    "💾 Atualizar status",
                    key=f"u{demanda.id}"
                ):
                    with SessionLocal() as db:
                        db.query(Demanda).filter(
                            Demanda.id == demanda.id
                        ).update(
                            {"status": novo}
                        )

                        db.commit()

                    st.rerun()

                if st.button(
                    "🗑️ Excluir demanda",
                    key=f"ed{demanda.id}"
                ):
                    excluir(
                        Demanda,
                        demanda.id
                    )

                    st.rerun()
elif pagina == "Kanban":
    st.title("📌 Kanban de Demandas")

    todas_demandas = demandas()

    colunas = st.columns(4)

    for i, status in enumerate(STATUS):
        with colunas[i]:
            demandas_coluna = [
                (demanda, cliente)
                for demanda, cliente in todas_demandas
                if demanda.status == status
            ]

            st.subheader(
                f"{status} ({len(demandas_coluna)})"
            )

            if not demandas_coluna:
                st.info("Nenhuma demanda")

            for demanda, cliente in demandas_coluna:
                nome_cliente = (
                    cliente.nome
                    if cliente
                    else "Cliente não encontrado"
                )

                with st.container(border=True):
                    st.write(
                        f"**{demanda.titulo}**"
                    )

                    st.caption(
                        f"🏛️ {nome_cliente}"
                    )

                    st.caption(
                        f"👤 {demanda.responsavel}"
                    )

                    st.caption(
                        f"🔥 Prioridade: {demanda.prioridade}"
                    )

                    if demanda.prazo:
                        st.caption(
                            f"📅 Prazo: "
                            f"{demanda.prazo.strftime('%d/%m/%Y')}"
                        )

                    novo_status = st.selectbox(
                        "Mover para",
                        STATUS,
                        index=STATUS.index(demanda.status),
                        key=f"kanban_status_{demanda.id}"
                    )

                    if novo_status != demanda.status:
                        with SessionLocal() as db:
                            db.query(Demanda).filter(
                                Demanda.id == demanda.id
                            ).update(
                                {"status": novo_status}
                            )

                            db.commit()

                        st.rerun()


elif pagina == "Carteiras":
    st.title("🗂️ Carteiras")

    st.subheader("📥 Importar carteira 2026")

    st.write(
        "Use o arquivo Excel da divisão de carteiras para cadastrar "
        "automaticamente os clientes e seus respectivos responsáveis."
    )

    arquivo = st.file_uploader(
        "Selecione a planilha da carteira",
        type=["xlsx"],
    )

    if arquivo is not None:
        st.success(
            f"Arquivo selecionado: {arquivo.name}"
        )

        if st.button(
            "🚀 Importar carteira 2026",
            use_container_width=True
        ):
            importados, ignorados = importar_carteira_excel(
                arquivo
            )

            st.success(
                f"Importação concluída! "
                f"{importados} cliente(s) importado(s)."
            )

            if ignorados > 0:
                st.info(
                    f"{ignorados} cliente(s) já existia(m) "
                    "no sistema e foi(ram) atualizado(s)."
                )

            st.rerun()

    st.divider()

    st.subheader("👥 Carteiras atuais")

    cs = clientes()

    for responsavel in RESPONSAVEIS[:3]:
        grupo = [
            c
            for c in cs
            if c.responsavel == responsavel
        ]

        with st.expander(
            f"{responsavel} — {len(grupo)} cliente(s)"
        ):
            if not grupo:
                st.info(
                    "Nenhum cliente cadastrado."
                )

            for cliente in grupo:
                st.write(
                    f"• **{cliente.nome}** — {cliente.tipo}"
                )


elif pagina == "Relatórios":
    st.title("📈 Relatórios")

    st.selectbox(
        "Tipo de relatório",
        [
            "Relatório mensal",
            "Relatório de transparência",
            "Relatório de demandas",
            "Relatório de produtividade"
        ]
    )

    st.info(
        "A geração de relatórios será implementada "
        "em uma próxima etapa."
    )


else:
    st.title("⚙️ Configurações")

    st.checkbox(
        "Ativar notificações internas"
    )

    st.checkbox(
        "Mostrar alertas de demandas atrasadas"
    )

    st.write(
        "**Sistema:** Facilita Manager"
    )

    st.write(
        "**Tecnologia:** Python + Streamlit + SQLite"
    )

    st.write(
        "**Versão:** 0.2.0"
    )